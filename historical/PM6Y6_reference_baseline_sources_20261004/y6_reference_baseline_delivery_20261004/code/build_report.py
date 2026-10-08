"""Build the Chinese PM6:Y6 baseline report from saved, audited results only.
Run with $CODEX_PRIMARY_RUNTIME_PYTHON. Does not run or modify the solver/data.
"""
from pathlib import Path
import json, os, subprocess, math
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image
ROOT=Path(__file__).resolve().parents[1]
QA=ROOT/'report_qa'; QA.mkdir(exist_ok=True)
OUT=ROOT/'PM6Y6_reference_baseline_report.docx'
cal=json.loads((ROOT/'data/calibration_audit.json').read_text())
obs=json.loads((ROOT/'data/observable_validation.json').read_text())
metrics={int(d['T']):d for p in (ROOT/'data').glob('*_metrics.json') for d in [json.loads(p.read_text())]}

# Raster equations via bundled Matplotlib; MathJax is absent in this runtime.
equations={
 'fd':r'c(\eta,s)=\frac{n}{N_0}=\int_{-\infty}^{\infty}\frac{e^{-z^2/2}}{\sqrt{2\pi}}\,\frac{dz}{1+e^{sz-\eta}},\quad s=\frac{\sigma}{k_{\mathrm{B}}T}',
 'einstein':r'\frac{D}{\mu}=\frac{k_{\mathrm{B}}T}{q}\,\frac{c}{\partial c/\partial\eta}',
 'mobility':r'\frac{\mu_i(T)}{\mu_i(300)}=\exp\!\left[-\frac{4}{9}\left(\frac{\sigma_i}{k_{\mathrm{B}}}\right)^2\left(\frac{1}{T^2}-\frac{1}{300^2}\right)\right]',
 'reaction':r'R=\beta(T)np\left[1-e^{-A}\right],\quad A=\frac{E_{\mathrm{Fn}}-E_{\mathrm{Fp}}}{k_{\mathrm{B}}T}',
 'beta':r'\beta(T)=\beta_{300}\left[\frac{n_{\mathrm{PIA}}(300)}{n_{\mathrm{PIA}}(T)}\right]^2,\quad G_0=\beta_{300}n_{\mathrm{PIA}}(300)^2'
}
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
for name,tex in equations.items():
    fig=plt.figure(figsize=(9,1.0))
    fig.text(.5,.5,'$'+tex+'$',ha='center',va='center',fontsize=18)
    fig.savefig(QA/(name+'.png'),dpi=300,bbox_inches='tight',pad_inches=.035,transparent=True)
    plt.close(fig)

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.70);sec.right_margin=Inches(.70)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.25)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=doc.styles[name];st.font.name='Noto Serif CJK SC';st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
 st.element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Calibri')
 st.element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Calibri')
 st.paragraph_format.space_after=Pt(6)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.12
doc.styles['Title'].font.size=Pt(24);doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(10)
doc.styles['Subtitle'].font.size=Pt(10)
doc.styles['Heading 1'].font.size=Pt(16);doc.styles['Heading 1'].font.bold=True
doc.styles['Heading 1'].paragraph_format.space_before=Pt(0);doc.styles['Heading 1'].paragraph_format.space_after=Pt(9)
doc.styles['Heading 2'].font.size=Pt(11.5);doc.styles['Heading 2'].font.bold=True
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10);doc.styles['Heading 2'].paragraph_format.space_after=Pt(5)
doc.styles['Caption'].font.size=Pt(8.5);doc.styles['Caption'].font.italic=False
for st in ['Title','Subtitle','Heading 1','Heading 2']:
 el=doc.styles[st].element
 for b in el.xpath('.//w:pBdr'): b.getparent().remove(b)
header=sec.header.paragraphs[0];header.text='PM6:Y6 变温参考基线  |  条件性模型与验证审计';header.style=doc.styles['Caption']
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('2026年10月4日  ·  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

def p(text='',boldlead=None,style=None):
 para=doc.add_paragraph(style=style)
 if boldlead and text.startswith(boldlead):
  para.add_run(boldlead).bold=True;para.add_run(text[len(boldlead):])
 else:para.add_run(text)
 para.paragraph_format.widow_control=True
 return para

pending_page=False
def h(text,level=1):
 global pending_page
 para=doc.add_paragraph(text,style=f'Heading {level}')
 if pending_page:para.paragraph_format.page_break_before=True;pending_page=False
 return para
def page():
 global pending_page
 pending_page=True
def caption(text):return p(text,style='Caption')
def picture(path,max_w=7.1,max_h=5.4,alt=''):
 path=Path(path);w0,h0=Image.open(path).size;w=min(max_w,max_h*w0/h0)
 para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
 para.paragraph_format.space_after=Pt(3)
 im=para.add_run().add_picture(str(path),width=Inches(w))
 im._inline.docPr.set('descr',alt or path.stem)
 return para

def eq(name,max_w=6.6,max_h=.48):picture(QA/(name+'.png'),max_w,max_h,'数学表达式 '+name)
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for j,txt in enumerate(headers):t.rows[0].cells[j].text=txt
 for row in rows:
  cells=t.add_row().cells
  for j,txt in enumerate(row):cells[j].text=str(txt)
 for i,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr()
  cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if i==0:
   rep=OxmlElement('w:tblHeader');trPr.append(rep)
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcPr=c._tc.get_or_add_tcPr()
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DEE7EC' if i==0 else ('F6F8F9' if i%2==0 else 'FFFFFF'));tcPr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   tcPr.append(borders)
   mar=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
   tcPr.append(mar)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1);para.paragraph_format.line_spacing=1.05
    para.alignment=WD_ALIGN_PARAGRAPH.CENTER if j==0 or (len(headers)>3 and j>0) else WD_ALIGN_PARAGRAPH.LEFT
    for r in para.runs:r.font.size=Pt(9);r.bold=(i==0)
 p().paragraph_format.space_after=Pt(1)
 return t

# PAGE 1
p('PM6 Y6 变温参考基线修正报告',style='Title')
p('100–300 K 高斯态密度器件模型  |  实现检查与物理验证  |  2026年10月4日',style='Subtitle')
p('结论：数值闭合已经修正，完整实验验证仍未通过。',boldlead='结论：')
p('本报告核对 PM6:Y6 文献参数、费米–狄拉克统计、电子与空穴非对称迁移率、条件性复合系数，以及一维漂移扩散器件解。结果可作为可追溯的参考基线与失败诊断，不构成全温区实验 J–V 拟合，也不能据此认定低温输运机制已被确定。')
p('最明显的否定证据在 100 K：条件性器件解的短路电流仅为 3.486×10⁻⁶ mA cm⁻²，开路电压为 1.10463 V；文献 Fig. 5a 的开路电压约为 0.960 V，而 Fig. S15a 在 0.04 V 的光照电流幅值约为 17.9 mA cm⁻²。电压和低偏压电流均出现重大偏离。[1–3]')
h('工作状态',2)
table(['项目','状态','判定依据'],[
 ['高斯 DOS 与精确 FD','已完成','独立积分、反函数与导数检查通过'],
 ['非对称迁移率与符号修正','已完成','300 K 锚点分别为 8.4×10⁻⁴ 与 1.3×10⁻⁴'],
 ['复合 β 与光生率 G','条件性','PIA 密度与恒定 G 假设共同定义'],
 ['接触边界与介电常数','条件性','零势垒理想库、少数载流子阻挡；εᵣ=3.5'],
 ['PDE 光照与数值收敛','已完成','已求解 Jsc、Voc 与 MPP；并作选点细网格检查'],
 ['100–300 K 完整实验匹配','验证未通过','低温电流、电压与独立复合趋势未共同吻合']
 ],[1.95,.85,4.3])
h('读者应保留的边界',2)
p('PIA 的 n(T) 是校准输入；由它反推出 β(T) 后，不能再把该 n(T) 的再现计作独立预测。SCLC 迁移率仅在约 223–328 K 有测量支持，低于 223 K 的所有器件结果必须标注外推。没有新增陷阱或 Marcus 通道，也没有逐温度调整带隙以抹平残差。')

# PAGE 2
page();h('1 参数来源与实验条件')
p('采用 Perdigón-Toro 等人 2022 年论文及其补充材料作为参考。[1] 表中严格区分实测或发表数值、作者拟合值、推断量与本次模型假设。')
table(['参数','数值','证据与使用限制'],[
 ['σₑ 与 σₕ','60 与 74 meV','正文 Table 1；两个高斯 DOS 宽度'],
 ['μₑ 与 μₕ 在 300 K','8.4×10⁻⁴ 与 1.3×10⁻⁴\ncm² V⁻¹ s⁻¹','SCLC 零场有效迁移率；非任意密度或电场测量'],
 ['N₀','2.4×10²⁰ cm⁻³','作者固定的分子数密度；非独立传输晶格间距'],
 ['活性层厚度','110 nm','SI 约值；正文另称 100 nm；逐器件厚度未知'],
 ['β₃₀₀','约 8×10⁻¹² cm³ s⁻¹','正文 Fig. 1 附近的近似值；样品条件需保留'],
 ['文献 DOS 能隙','1.43 与 1.41 eV','Table S3 两个解析区间的拟合值；非独立测量'],
 ['本次全局热力学 Eg','1.407221 eV','仅拟合 Fig. 5a 单一常数；RMSE 19.35 mV'],
 ['器件 PDE 的 Eg','1.42 eV','固定诊断输入；没有使用逐温度能隙'],
 ['εᵣ 与接触势垒','3.5；bₑ=bₕ=0','本次理想化假设，未由论文辨识']
 ],[1.65,2.0,3.45])
h('校准集与验证集',2)
p('Fig. 5b 给出开路 PIA 密度，按作者解释取 n=p，而非 n+p；Fig. 5a 允许拟合一个全局能量偏移。Fig. S16a 的 27 个低光强点（0.5、0.25、0.08 sun）保留为验证集；其 1 sun 点与 Fig. 5 重复，不计作独立验证。Fig. S15a 仅明确标示 300 K 和 100 K，不能给中间曲线擅配温度。')
p('条件不等同：PIA 使用 405 nm 连续光，BACE 使用 445 nm 调制激光，变温 J–V 使用 1 sun 等效白光；器件面积与银电极配置也存在区别。同一材料或名义结构不保证同一吸收光生率、同一器件。缺少 PM6:Y6 的绝对 PIA 截面及其温度依赖。[2]')
p('原文校正：SI S1c 的迁移率指数原 PDF 缺负号；这是源文问题，并非 MinerU 漏字。本基线采用标准负指数并明确记录此模型选择。OCR 还漏去 Table S3 的拟合参数脚注，解释时已恢复。[3]',boldlead='原文校正：')

# PAGE 3
page();h('2 统计 输运与复合闭合')
h('精确 FD 与广义 Einstein 关系',2)
p('分别对电子和空穴的高斯 DOS 积分，保留简并占据与正的密度响应。η 表示相对相应 DOS 中心的约化化学势，c 为占据分数；下式以电子记号示意，空穴同理。')
eq('fd',6.8,.5);eq('einstein',4.2,.39)
p('采用有限体积 Poisson 与连续性方程，以及基于割线逆可压缩性的 Scharfetter–Gummel 面通量。电势以 V、能量以 eV、长度以 cm 表示；代码中 kBT/q 的伏特数值与 kBT 的 eV 数值一致。面通量保持严格平衡和耗散符号，并在稀薄极限回到经典关系。[4]')
h('迁移率明确分开标定',2)
eq('mobility',6.0,.5)
p('电子、空穴分别使用各自的 300 K 锚点与 σ。该式是有效零场 SCLC 温度律，不能直接称为微观零密度迁移率。100 K 外推得到 μₑ≈4.04×10⁻¹²、μₕ≈2.90×10⁻¹⁷ cm² V⁻¹ s⁻¹；空穴输运急剧变慢是此基线电流坍塌的重要组成。未经新数据校准，不叠加 EGDM 场或密度增强因子。')
h('可逆暗态反应与条件性光生率',2)
eq('reaction',6.0,.43);eq('beta',6.8,.45)
p('反应在准费米能级相等时归零，并满足 R·A≥0。外加光生率 G 是非平衡源项；电化学功恒等式通过，不等于已建立完整光子能量或总熵账本。')
p(f'条件性归一化给出 G₀=1.27468×10²² cm⁻³ s⁻¹，110 nm 下 qG₀d={cal["Jgeneration_110nm_mAcm2"]:.3f} mA cm⁻²。它低于另一 AM1.5 条件下报道的 Jsc=24.9 mA cm⁻²，因此不能将两种条件直接等同，也不报告伪精确的 AM1.5 效率。')
p('接触边界：bₑ=bₕ=0 对应多数载流子接触处半 DOS 占据，是未经验证的高注入理想库；少数载流子阻挡与平衡库只是灵敏度场景。接触参数未辨识之前，不能把完整器件输出当作材料本征预言。',boldlead='接触边界：')

# PAGE 4
page();h('3 校准残差与独立留出验证')
p('一个固定 Eg 无法让全部温度点同时贴合。精确 FD 热力学校准 RMSE 为 19.35 mV；S16 低光强留出点的绝对 Voc RMSE 为 19.00 mV，去除同温 1 sun 基线偏移后的光强位移 RMSE 为 9.79 mV。后者不依赖 Eg，仍暴露光强响应误差。')
picture(ROOT/'figures/observable_validation.png',7.1,5.17,'四面板比较热力学校准、S16光强位移、BACE复合趋势和迁移率外推')
caption('图 1  校准与验证必须分开。左上使用 PIA 输入；右上为留出光强位移；左下对比不同实验条件的复合趋势；右下阴影为无 SCLC 测量支持的低温外推。误差条仅为读图界限。')
p('独立复合趋势未通过一致性检验。Fig. 3a 的 300 K 有效 R/n² 中位数约 2.10×10⁻¹¹ cm³ s⁻¹，为正文近似锚点的 2.63 倍；200 K 归一化 BACE 比值为 0.0458，恒定 G 的 PIA 推断为 0.1445，相差约 3.15 倍。应保留跨条件不确定性，不能任意重标定数据或宣称论文提取有误。')
p('读图不确定性不是实验置信区间。当前界限为 Fig. 5a ±5 mV、S16 ±6 mV、PIA ±0.025 decade、BACE ±0.10 decade、S15 电流 ±0.07 decade 与电压 ±6 mV；未报告的系统校准误差仍需额外计入。[2]')

# PAGE 5
page();h('4 条件性器件结果与失败证据')
p('下表来自 N=321 的收敛 PDE 状态，Voc 由零电流根求解，MPP 由再求解后的功率极值确定，非稀疏曲线插值代替。所有温度均取恒定 G、固定 Eg=1.42 eV 与理想阻挡接触。Jsc 按输出电流幅值列出。')
rows=[]
for T in [300,250,200,150,100]:
 d=metrics[T];j=abs(d['Jsc_Acm2'])*1000
 rows.append([str(T),f'{j:.4f}' if T>100 else '3.486×10⁻⁶',f'{d["Voc_V"]:.6f}',f'{d["FF"]*100:.2f}',f'{d["Vmp_V"]:.4f}'])
table(['T  K','Jsc  mA cm⁻²','Voc  V','FF  %','Vmpp  V'],rows,[.7,1.9,1.6,1.2,1.7])
fig=ROOT/'figures/conditional_jv.png'
if fig.exists():
 picture(fig,7.1,3.65,'条件性PM6:Y6温度相关JV曲线，低温外推与参考端点比较')
 caption('图 2  条件性 J–V 与温度响应。曲线仅表示当前参数组合的解，不能将画线平滑度解释为实验验证。')
else:
 p('条件性 J–V 图尚未纳入此构建；数值指标见上表。')
p('300 K 的 21.994 mA cm⁻²、0.799637 V、64.06% 只是条件性结果。文献报告的 FF 为 66.8%，并非 78%；不同光照归一化下的电流差异也不应靠任意效率目标掩盖。')
p('100 K 是决定性的反例：模型 Jsc 极小，而 S15 的低偏压电流幅值仍约 17.9 mA cm⁻²；PDE Voc 比 Fig. 5a 高约 144 mV。两者电压点并不完全相同，S15 也只有电流幅值，但相差约七个数量级的低偏压输运无法由读图误差解释。')
coverage_file=ROOT/'report_qa/final_coverage.txt'
p(coverage_file.read_text().strip() if coverage_file.exists() else '运行覆盖：本报告确认 100–300 K 的上述光照指标；100 K 暗态延拓仍在检查，不能声称已经验证完整 0–1.1 V 暗态覆盖。')

# PAGE 6
page();h('5 数值可信度与下一步判据')
p('数值验证回答“方程是否被一致求解”，物理验证回答“这些方程和参数是否描述实验”。本次前者有独立检查支持，后者在低温明确失败。')
table(['检查','结果','适用范围'],[
 ['FD 独立积分','密度相对误差 ≤1.16×10⁻¹²\n导数相对误差 ≤4.41×10⁻¹⁰','117 项检查；包含 100 K 与两种 DOS 宽度'],
 ['平衡与离散热力学','80 项面通量检查；54 项器件检查','零平衡电流、通量与反应耗散符号、功恒等式'],
 ['N 321 → 641 的 Jsc','300 K  0.00259%\n200 K  0.02165%\n100 K  0.2197%','选点空间离散差异；不含材料参数误差'],
 ['Voc 精细零点包围','端点电流符号均已分辨\n包围宽度约 6.10 μV','三温度网格差异均低于约 10 μV 报告分辨率']
 ],[1.55,2.65,2.9])
p('极小电流必须按终端电流本身检查。采用电流散布、局部与积分连续性误差及浮点相消量形成闭合指标，并施加 0.1% 相对电流目标和机器精度余量。开路处用局部 dJ/dV 将残差换算为电压，并要求有可分辨符号的包围区间；仅有绝对残差小，不能证明低温电流可信。[4]')
h('优先补齐的物理输入',2)
p('1  同器件、同光谱条件的 G(T,F) 与绝对 PIA 标定，先解除光生率和复合归一化的混淆。\n2  低于 223 K 的迁移率证据，以及实际工作密度和电场下的输运；避免把外推温度律继续当作测量。\n3  接触功函数、注入势垒、表面复合速度与独立厚度；用暗态和光照曲线联合约束边界。\n4  使用留出光强、BACE 与完整带符号 J–V 再验证。新增机制前先报告现有基线的失配，不用新陷阱或逐温参数掩盖失败。')
h('来源与复现索引',2)
caption('[1] Perdigón-Toro et al. Understanding the Role of Order in Y-Series Non-Fullerene Solar Cells to Realize High Open-Circuit Voltages. Advanced Energy Materials, 2022. https://doi.org/10.1002/aenm.202103422；正文 Table 1、Fig. 1/3/5；SI S1c、Table S3、Fig. S12/S15/S16。')
caption('[2] 本包 reference/conditions_and_uncertainty.json 与 parameter_ledger.csv；图像中心数字化数据保留了独立的读图界限及跨条件说明。')
caption('[3] reference/SOURCE_OCR_QA.md；原 PDF 与 MinerU 原始输出分开保存，人工核对不伪装成 OCR。')
caption('[4] independent_review/PHYSICAL_REVIEW.md、fd_eos_results.json、mesh_refine_T*.json、voc_refine_T*.json；data/calibration_audit.json、observable_validation.json 与 *_metrics.json。本报告由 code/build_report.py 生成。')

doc.core_properties.title='PM6 Y6 变温参考基线修正报告'
doc.core_properties.subject='条件性器件模拟的数值闭合与否定性物理验证'
doc.core_properties.author=''
doc.core_properties.keywords='PM6:Y6, Gaussian DOS, Fermi Dirac, drift diffusion, validation'
doc.save(OUT)
print(OUT)
